"""Excel report of follow-up outcomes: what was done on time, late or missed, and what the calls achieved.

The "Follow-ups" and "Calls" tabs hold one row per record; every number on "Summary" and "Daily trend"
is a COUNTIFS formula over those rows, so the workbook stays correct if someone filters or edits it.
"""
from __future__ import annotations

import json

from datetime import date, datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .sarthi import target as sarthi_target
from .context import Ctx, iso, parse_dt
from .db import rows
from .followups import KIND_LABELS, OUTCOMES
from .heat import TIER_BADGE
from .reports import classify

FONT = "Arial"
# status colours (good / warning / serious / critical), always paired with a text label
GOOD, WARN, SERIOUS, CRIT = "0CA30C", "FAB219", "EC835A", "D03B3B"
FILLS = {"On time": "DDF3DD", "Late": "FEF0CC", "Overdue": "FBE3D9", "Missed": "F6D8D8", "Upcoming": "EEEEEE"}
RESULT_LABEL = {"on_time": "On time", "late": "Late", "missed": "Missed", "overdue": "Overdue", "upcoming": "Upcoming"}
REACHED = {"interested", "connected", "callback_requested", "converted", "referral_given"}
NEGATIVE = {"not_interested", "wrong_number"}
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
THIN = Border(bottom=Side(style="thin", color="DDDDDD"))


def call_result(outcome: str | None) -> str:
    if not outcome:
        return ""
    if outcome == "converted":
        return "Converted"
    if outcome in REACHED:
        return "Reached"
    return "Closed (negative)" if outcome in NEGATIVE else "Not reached"


def _font(bold=False, color="222222", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _header(ws, row: int, values: list[str], widths: list[int] | None = None) -> None:
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 30


def _write(ws, row: int, values: list, fmt: dict[int, str] | None = None) -> None:
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = _font()
        c.border = THIN
        if fmt and i in fmt:
            c.number_format = fmt[i]


def collect(ctx: Ctx, start: date, end: date) -> tuple[list[dict], list[dict]]:
    """Follow-ups due and calls made between start and end (inclusive)."""
    lo = iso(datetime.combine(start, time(0), ctx.tz))
    hi = iso(datetime.combine(end + timedelta(days=1), time(0), ctx.tz))
    now = ctx.now()
    grace = timedelta(minutes=ctx.rules.get("warning_after_minutes", 15))
    fus = []
    for f in rows(ctx.conn, """SELECT f.*, l.name, l.category, l.heat_tier, c.outcome, c.notes
                              FROM followups f JOIN leads l ON l.lead_id=f.lead_id
                              LEFT JOIN calls c ON c.id=f.call_id
                              WHERE f.due_at>=? AND f.due_at<? AND f.status!='cancelled' ORDER BY f.due_at""", (lo, hi)):
        due = parse_dt(f["due_at"], ctx.tz)
        done = parse_dt(f["completed_at"], ctx.tz)
        gm = ctx.member(f["agent_id"]) or {}
        late = max(int(((done or now) - due).total_seconds() // 60), 0) if (done or now) > due + grace else 0
        fus.append({
            "due": due.replace(tzinfo=None), "gm": gm.get("name", f["agent_id"]),
            "sm": (ctx.member(ctx.sm_of(gm)) or {}).get("name", ""), "client": f["lead_id"], "name": f["name"] or "",
            "priority": TIER_BADGE.get(f["heat_tier"] or "cold"), "category": f["category"] or "",
            "type": KIND_LABELS.get(f["kind"], f["kind"]), "reason": f["reason"] or "",
            "result": RESULT_LABEL[classify(ctx, f, now)], "late_min": late or None,
            "done": done.replace(tzinfo=None) if done else None,
            "outcome": OUTCOMES.get(f["outcome"] or "", {}).get("label", ""), "call_result": call_result(f["outcome"]),
            "escalated": "Yes" if f["missed"] else "", "notes": f["notes"] or "",
        })
    calls = []
    for c in rows(ctx.conn, """SELECT c.*, l.name, l.category, l.heat_tier, l.extra FROM calls c JOIN leads l ON l.lead_id=c.lead_id
                              WHERE c.called_at>=? AND c.called_at<? AND c.program IS NULL AND COALESCE(l.campaign,'') NOT IN ('referral','sarthi')
                              ORDER BY c.called_at""", (lo, hi)):
        gm = ctx.member(c["agent_id"]) or {}
        calls.append({
            "at": parse_dt(c["called_at"], ctx.tz).replace(tzinfo=None), "gm": gm.get("name", c["agent_id"]),
            "client": c["lead_id"], "name": c["name"] or "", "priority": TIER_BADGE.get(c["heat_tier"] or "cold"),
            "category": c["category"] or "", "outcome": OUTCOMES[c["outcome"]]["label"],
            "result": call_result(c["outcome"]), "notes": c["notes"] or "",
            "sarthi": "Yes" if c.get("sarthi_activated") else "",
            "deposit": c.get("deposit_amount") or None,
            "funded": "Yes" if c["outcome"] == "converted" and sarthi_target(ctx) > 0 and max(
                c.get("deposit_amount") or 0, json.loads(c.get("extra") or "{}").get("current_margin") or 0) >= sarthi_target(ctx) else "",
            "next": parse_dt(c["next_followup_at"], ctx.tz).replace(tzinfo=None) if c["next_followup_at"] else None,
        })
    return fus, calls


def build_workbook(ctx: Ctx, start: date, end: date, title_note: str = "") -> Workbook:
    fus, calls = collect(ctx, start, end)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    fu_ws = wb.create_sheet("Follow-ups")
    call_ws = wb.create_sheet("Calls")
    trend = wb.create_sheet("Daily trend")
    notes = wb.create_sheet("How to read")

    # ---------------------------------------------------------------- detail: follow-ups
    fu_cols = ["Due", "GM", "SM", "Client code", "Name", "Priority", "Category", "Follow-up type", "Reason",
               "Result", "Minutes late", "Done at", "Call outcome", "Call result", "Escalated to SM", "Notes"]
    _header(fu_ws, 1, fu_cols, [17, 16, 16, 13, 16, 10, 26, 16, 34, 11, 9, 17, 26, 15, 10, 34])
    for r, f in enumerate(fus, 2):
        _write(fu_ws, r, [f["due"], f["gm"], f["sm"], f["client"], f["name"], f["priority"], f["category"], f["type"],
                          f["reason"], f["result"], f["late_min"], f["done"], f["outcome"], f["call_result"],
                          f["escalated"], f["notes"]], {1: "dd mmm yyyy hh:mm", 12: "dd mmm yyyy hh:mm", 11: "0"})
    n_fu = max(len(fus) + 1, 2)
    fu_ws.freeze_panes = "B2"
    fu_ws.auto_filter.ref = f"A1:{get_column_letter(len(fu_cols))}{n_fu}"
    for label, color in FILLS.items():
        fu_ws.conditional_formatting.add(f"A2:{get_column_letter(len(fu_cols))}{n_fu}",
                                         FormulaRule(formula=[f'$J2="{label}"'], fill=PatternFill("solid", fgColor=color)))

    # ---------------------------------------------------------------- detail: calls
    call_cols = ["Called at", "GM", "Client code", "Name", "Priority", "Category", "Outcome", "Call result", "Notes",
                 "Next follow-up", "Sarthi activated", "Deposit (₹)", "Funded to target"]
    _header(call_ws, 1, call_cols, [17, 16, 13, 16, 10, 26, 28, 16, 40, 17, 10, 12, 10])
    for r, c in enumerate(calls, 2):
        _write(call_ws, r, [c["at"], c["gm"], c["client"], c["name"], c["priority"], c["category"], c["outcome"],
                            c["result"], c["notes"], c["next"], c["sarthi"], c["deposit"], c["funded"]],
               {1: "dd mmm yyyy hh:mm", 10: "dd mmm yyyy hh:mm", 12: "#,##0"})
    n_call = max(len(calls) + 1, 2)
    call_ws.freeze_panes = "B2"
    call_ws.auto_filter.ref = f"A1:{get_column_letter(len(call_cols))}{n_call}"
    for label, color in {"Converted": "C9EBC9", "Reached": "DDF3DD", "Not reached": "FEF0CC", "Closed (negative)": "F6D8D8"}.items():
        call_ws.conditional_formatting.add(f"A2:{get_column_letter(len(call_cols))}{n_call}",
                                           FormulaRule(formula=[f'$H2="{label}"'], fill=PatternFill("solid", fgColor=color)))

    F = lambda col: f"'Follow-ups'!${col}$2:${col}${n_fu}"  # noqa: E731
    C = lambda col: f"Calls!${col}$2:${col}${n_call}"  # noqa: E731

    # ---------------------------------------------------------------- summary
    ws = summary
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 13
    ws["A1"] = "Follow-up outcomes report"
    ws["A1"].font = _font(bold=True, size=16, color="1F3A5F")
    ws["A2"] = f"{start:%d %b %Y} to {end:%d %b %Y}" + (f"  ·  {title_note}" if title_note else "")
    ws["A2"].font = _font(italic=True, color="666666")
    if title_note:
        ws["A2"].font = _font(bold=True, color=CRIT)

    kpis = [
        ("Follow-ups due", f"=COUNTIFS({F('J')},\"<>Upcoming\")", "0"),
        ("✅ Done on time (success)", f"=COUNTIFS({F('J')},\"On time\")", "0"),
        ("⚠️ Done late", f"=COUNTIFS({F('J')},\"Late\")", "0"),
        ("❌ Missed (escalated to SM)", f"=COUNTIFS({F('J')},\"Missed\")", "0"),
        ("❌ Still overdue", f"=COUNTIFS({F('J')},\"Overdue\")", "0"),
        ("On-time rate", "=IFERROR(B6/B5,0)", "0.0%"),
        ("Fail rate (missed + overdue)", "=IFERROR((B8+B9)/B5,0)", "0.0%"),
        ("Calls made", f"=COUNTIFS({C('A')},\"<>\")", "0"),
        ("Customers reached", f"=COUNTIFS({C('H')},\"Reached\")+COUNTIFS({C('H')},\"Converted\")", "0"),
        ("Not reached (NPC / busy / off)", f"=COUNTIFS({C('H')},\"Not reached\")", "0"),
        ("Converted", f"=COUNTIFS({C('H')},\"Converted\")", "0"),
        ("Reach rate", "=IFERROR(B13/B12,0)", "0.0%"),
        ("Sarthi activated on calls", f"=COUNTIFS({C('K')},\"Yes\")", "0"),
        (f"Funded ≥ ₹{sarthi_target(ctx):,} (OKR)",
         f"=COUNTIFS({C('M')},\"Yes\")", "0"),
        ("Deposits on converted calls (₹)", f"=SUMIFS({C('L')},{C('H')},\"Converted\")", "#,##0"),
    ]
    ws["A4"], ws["B4"] = "Headline", "Value"
    for c in (ws["A4"], ws["B4"]):
        c.font, c.fill = _font(bold=True, color="FFFFFF"), HEAD_FILL
    for i, (label, formula, fmt) in enumerate(kpis, 5):
        ws.cell(row=i, column=1, value=label).font = _font(bold=label.startswith(("On-time", "Fail", "Reach")))
        c = ws.cell(row=i, column=2, value=formula)
        c.font, c.number_format = _font(bold=True), fmt
    ws.conditional_formatting.add("B10", CellIsRule(operator="greaterThanOrEqual", formula=["0.9"], font=Font(name=FONT, bold=True, color=GOOD)))
    ws.conditional_formatting.add("B10", CellIsRule(operator="lessThan", formula=["0.7"], font=Font(name=FONT, bold=True, color=CRIT)))
    ws.conditional_formatting.add("B11", CellIsRule(operator="greaterThan", formula=["0.1"], font=Font(name=FONT, bold=True, color=CRIT)))

    def table(top: int, title: str, key_col: str, keys: list[str], first: str) -> int:
        ws.cell(row=top, column=1, value=title).font = _font(bold=True, size=12, color="1F3A5F")
        heads = [first, "Due", "On time", "Late", "Missed", "Overdue", "On-time %", "Calls", "Reached", "Converted"]
        for i, h in enumerate(heads, 1):
            c = ws.cell(row=top + 1, column=i, value=h)
            c.font, c.fill, c.alignment = _font(bold=True, color="FFFFFF"), HEAD_FILL, Alignment(wrap_text=True)
        call_key = {"B": "B", "F": "E", "H": None, "G": "F"}.get(key_col)  # same dimension on the Calls tab
        for r, k in enumerate(keys, top + 2):
            ws.cell(row=r, column=1, value=k).font = _font()
            crit = f"{F(key_col)},$A{r}"
            vals = [f"=COUNTIFS({crit},{F('J')},\"<>Upcoming\")",
                    f"=COUNTIFS({crit},{F('J')},\"On time\")", f"=COUNTIFS({crit},{F('J')},\"Late\")",
                    f"=COUNTIFS({crit},{F('J')},\"Missed\")", f"=COUNTIFS({crit},{F('J')},\"Overdue\")",
                    f"=IFERROR(C{r}/B{r},0)"]
            if call_key:
                cc = f"{C(call_key)},$A{r}"
                vals += [f"=COUNTIFS({cc})", f"=COUNTIFS({cc},{C('H')},\"Reached\")+COUNTIFS({cc},{C('H')},\"Converted\")",
                         f"=COUNTIFS({cc},{C('H')},\"Converted\")"]
            for i, v in enumerate(vals, 2):
                c = ws.cell(row=r, column=i, value=v)
                c.font, c.border = _font(), THIN
                c.number_format = "0.0%" if i == 7 else "0"
        last = top + 1 + len(keys)
        if keys:
            rng = f"G{top + 2}:G{last}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["0.9"], fill=PatternFill("solid", fgColor="DDF3DD")))
            ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["0.7", "0.8999"], fill=PatternFill("solid", fgColor="FEF0CC")))
            ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0.7"], fill=PatternFill("solid", fgColor="F6D8D8")))
        return last + 2

    gms = [m["name"] for m in ctx.callers()]
    nxt = table(5 + len(kpis) + 1, "By GM (who is following up on time)", "B", gms, "GM")
    nxt = table(nxt, "By follow-up type", "H", list(KIND_LABELS.values()), "Follow-up type")
    nxt = table(nxt, "By priority", "F", list(TIER_BADGE.values()), "Priority")
    cats = sorted({f["category"] for f in fus} | {c["category"] for c in calls})
    table(nxt, "By customer category", "G", cats, "Category")
    ws.freeze_panes = "A4"

    # ---------------------------------------------------------------- daily trend + chart
    _header(trend, 1, ["Day", "On time", "Late", "Missed", "Overdue", "On-time %", "Calls", "Reached"], [14, 10, 10, 10, 10, 11, 10, 10])
    days, d = [], start
    while d <= end:
        if ctx.is_working_day(d):
            days.append(d)
        d += timedelta(days=1)
    for r, d in enumerate(days, 2):
        lo, hi = datetime.combine(d, time(0)), datetime.combine(d + timedelta(days=1), time(0))
        trend.cell(row=r, column=1, value=d).number_format = "ddd dd mmm"
        win = f"{F('A')},\">=\"&DATE({lo.year},{lo.month},{lo.day}),{F('A')},\"<\"&DATE({hi.year},{hi.month},{hi.day})"
        cwin = f"{C('A')},\">=\"&DATE({lo.year},{lo.month},{lo.day}),{C('A')},\"<\"&DATE({hi.year},{hi.month},{hi.day})"
        for i, res in enumerate(["On time", "Late", "Missed", "Overdue"], 2):
            trend.cell(row=r, column=i, value=f"=COUNTIFS({win},{F('J')},\"{res}\")")
        trend.cell(row=r, column=6, value=f"=IFERROR(B{r}/SUM(B{r}:E{r}),0)").number_format = "0.0%"
        trend.cell(row=r, column=7, value=f"=COUNTIFS({cwin})")
        trend.cell(row=r, column=8, value=f"=COUNTIFS({cwin},{C('H')},\"Reached\")+COUNTIFS({cwin},{C('H')},\"Converted\")")
        for col in range(1, 9):
            trend.cell(row=r, column=col).font = _font()
    if days:
        chart = BarChart()
        chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
        chart.title = "Follow-ups by result, per day"
        chart.y_axis.title = "Follow-ups"
        chart.add_data(Reference(trend, min_col=2, max_col=5, min_row=1, max_row=len(days) + 1), titles_from_data=True)
        chart.set_categories(Reference(trend, min_col=1, min_row=2, max_row=len(days) + 1))
        for series, color in zip(chart.series, [GOOD, WARN, CRIT, SERIOUS]):
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = "FFFFFF"
        chart.gapWidth = 60
        chart.height, chart.width = 9, 22
        trend.add_chart(chart, "J2")

    # ---------------------------------------------------------------- how to read
    grace, esc = ctx.rules.get("warning_after_minutes", 15), ctx.rules.get("escalate_after_minutes", 60)
    lines = [
        ("How to read this report", ""),
        ("", ""),
        ("Follow-up result", ""),
        ("On time (success)", f"Called within {grace} min of the due time"),
        ("Late", f"Called {grace}–{esc} min after the due time (warning sent to the GM)"),
        ("Missed (fail)", f"Not called within {esc} min: escalated to the SM. Counts as missed even if called later"),
        ("Overdue (fail)", "Past due and not yet called or escalated"),
        ("Upcoming", "Not due yet; left out of every rate"),
        ("", ""),
        ("Call result", ""),
        ("Reached", "Connected: interested, needs follow-up, or asked for a callback"),
        ("Converted", "Deposited / traded on the call"),
        ("Not reached", "Not picked up, busy or switched off: retried automatically after the NPC rule"),
        ("Closed (negative)", "Not interested or wrong number"),
        ("", ""),
        ("Rates", ""),
        ("On-time rate", "On time ÷ follow-ups due (upcoming excluded)"),
        ("Fail rate", "(Missed + overdue) ÷ follow-ups due"),
        ("Reach rate", "Customers reached (incl. converted) ÷ calls made"),
        ("Colours", "On-time % ≥ 90% green, 70–90% amber, below 70% red"),
        ("", ""),
        ("Source", f"Generated by the Growth follow-up system on {ctx.now():%d %b %Y %H:%M}. Every Summary and Daily trend "
                   "number is a formula over the Follow-ups and Calls tabs."),
    ]
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 100
    for r, (a, b) in enumerate(lines, 1):
        notes.cell(row=r, column=1, value=a).font = _font(bold=bool(a) and not b, size=14 if r == 1 else 10, color="1F3A5F" if not b else "222222")
        notes.cell(row=r, column=2, value=b).font = _font()
    summary["A5"].comment = Comment("Upcoming follow-ups (not due yet) are excluded from every count and rate.", "Growth follow-up system")
    return wb


def write_report(ctx: Ctx, start: date | None = None, end: date | None = None, path=None, title_note: str = "") -> str:
    end = end or ctx.today()
    start = start or end - timedelta(days=6)
    path = path or ctx.path(ctx.cfg.get("reports_dir", "data/reports")) / f"followup_outcomes_{start:%Y%m%d}_{end:%Y%m%d}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        build_workbook(ctx, start, end, title_note).save(path)
    except PermissionError:  # open in Excel
        path = path.with_name(f"{path.stem}_{ctx.now():%H%M}{path.suffix}")
        build_workbook(ctx, start, end, title_note).save(path)
    return str(path)
