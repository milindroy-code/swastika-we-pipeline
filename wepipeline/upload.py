"""Analyse a customer list uploaded by an SM: ingest, age from the client database, score, rank, assign, pitch."""
from __future__ import annotations

import json
import logging
from collections import Counter
from datetime import timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .context import Ctx, iso, norm_client_id
from .sarthi import money, target
from .db import one, rows
from .ingest import ingest_webengage_file

log = logging.getLogger(__name__)
# The analysed list a GM works from: who, what we know, why call, what to say. The rest is on the Details tab.
BATCH_COLUMNS = [
    ("Rank", "batch_rank"), ("Priority", "priority_score"), ("Client", "client_line"), ("Client profile", "profile"),
    ("App activity", "activity"), ("Engagement levers", "levers"), ("Insight", "script_insight"), ("Call script", "script_text"),
    ("Ask (call to action)", "script_ask"), ("Agar customer bole", "objections_hi"), ("When to call", "when"),
    ("GM", "gm_assigned"), ("Status", "status_line"),
]
WIDTHS = {"batch_rank": 7, "priority_score": 12, "client_line": 16, "profile": 34, "activity": 40, "levers": 40,
          "script_insight": 45, "script_text": 70, "script_ask": 50, "objections_hi": 48, "when": 18, "gm_assigned": 14,
          "status_line": 20}
DETAIL_COLUMNS = [
    ("Rank in this list", "batch_rank"), ("Overall rank", "rank"), ("Priority", "priority"), ("Heat score", "heat"),
    ("Client code", "client"), ("Name", "name"), ("Phone", "phone"), ("Category", "category"), ("Stage", "stage"),
    ("List", "list"), ("Assigned GM", "gm"), ("SM", "sm"), ("First call due", "next_due"), ("Registered on", "registered_on"),
    ("Days since registration", "days_since_registration"), ("Deposit status", "deposit_status"),
    ("Last trade date", "last_trade_date"), ("Core activity (type × times, recency)", "core_detail"),
    ("Non-core activity", "non_core_detail"), ("Best time to call", "best_time"), ("Research segments", "segments"), ("Sarthi status", "sarthi_status"),
    ("Why hot (score drivers)", "heat_reasons"), ("Engagement levers", "levers"),
    ("Call ideas (data → angle)", "call_ideas_text"), ("Close", "closing_text"), ("Agar customer bole", "objections_hi"), ("Client status", "status"),
]
# What the SM says the list is: sets the category (and so the pitch stage) for customers new to the system.
# The client database, when reachable, replaces it with the real ageing category.
LIST_TYPES = {
    "no_deposit": ("New clients without deposit (push to the {target})", "Not yet matched in client DB", "new"),
    "low_margin": ("New clients with low margin, deposits under {target} (top up to {target})", "Not yet matched in client DB", "low_margin"),
    "reactivation": ("Churned clients showing activity (resurrect; only engaged ones go to GMs)", "Not yet matched in client DB", "reactivation"),
    "payout": ("Clients who made payouts (keep them invested and trading)", "Not yet matched in client DB", "payout"),
    "sarthi_disabled": ("Switched Sarthi off in the last 15 days (win them back)", "Not yet matched in client DB", "sarthi_disabled"),
    "sarthi_off": ("New clients with Sarthi not switched on (switch it on together)", "Not yet matched in client DB", "sarthi_off"),
    "unknown": ("Mixed / not sure (work it out from app activity)", None, ""),  # not offered in the dashboard
}
TIER_FILL = {"🔥 Hot": "FCDDD9", "🌤️ Warm": "FFF2CC", "❄️ Cold": "E6EFFC"}


# A client is on one list only: when two lists hold them, the earlier one here wins. A recent Sarthi switch-off is urgent
# (just lost); Sarthi-not-on comes last, since a deposit or a top-up matters more than switching Sarthi on.
CAMPAIGN_PRECEDENCE = ["payout", "sarthi_disabled", "reactivation", "low_margin", "new", "sarthi_off"]


def analyse_upload(ctx: Ctx, path: Path, uploaded_by: str | None = None, use_client_db: bool = True,
                   list_type: str = "unknown") -> dict:
    from . import jobs
    from .clientdb import available, enrich_leads, reachable

    ingest = ingest_webengage_file(ctx, path, create_from_profiles=True, category=LIST_TYPES[list_type][1])
    ids = ingest["client_ids"]
    if not ids:
        return {"error": "No client codes found. The file needs a ClientCode / User ID column.", "ingest": ingest}
    known_before = ingest["users"] - ingest["new_leads"]
    campaign = LIST_TYPES[list_type][2]
    marks = ",".join("?" * len(ids))
    ctx.conn.execute(f"UPDATE leads SET status='open', assigned_to=NULL, assigned_at=NULL, npc_count=0, updated_at=? "
                     f"WHERE lead_id IN ({marks}) AND status='drive'", (iso(ctx.now()), *ids))  # drive-only clients: funding first
    if campaign:  # tag every customer in the file, new and already known, so the list can be filtered and handled
        if campaign in {"payout", "reactivation"}:
            # funded clients are exactly the ones who pay out or go dormant: bring them back into the follow-up flow
            ctx.conn.execute(f"UPDATE leads SET status='open', assigned_to=NULL, assigned_at=NULL, npc_count=0, updated_at=? "
                             f"WHERE lead_id IN ({marks}) AND status IN ('converted','out_of_window')", (iso(ctx.now()), *ids))
        # a client on several lists keeps the most urgent one (a payout is never downgraded by a later new-client file)
        rank = {c: i for i, c in enumerate(CAMPAIGN_PRECEDENCE)}
        for l in rows(ctx.conn, f"SELECT lead_id, campaign FROM leads WHERE lead_id IN ({marks}) AND status='open'", ids):
            if rank.get(campaign, 99) <= rank.get(l["campaign"] or "", 99):
                ctx.conn.execute("UPDATE leads SET campaign=? WHERE lead_id=?", (campaign, l["lead_id"]))
        ctx.conn.commit()
    released = 0
    fallback_stage = {"no_deposit": "kyc_no_deposit", "reactivation": "dormant_trader"}.get(list_type)
    if fallback_stage:
        for l in rows(ctx.conn, f"SELECT lead_id, extra FROM leads WHERE lead_id IN ({','.join('?' * len(ids))}) AND status='open'", ids):
            extra = json.loads(l["extra"] or "{}")
            if not extra.get("stage"):  # the client database, when reachable, replaces it with the real stage
                extra["stage"] = fallback_stage
                ctx.conn.execute("UPDATE leads SET extra=? WHERE lead_id=?", (json.dumps(extra), l["lead_id"]))
        ctx.conn.commit()
    if campaign == "payout":
        _store_payout_details(ctx, path, ids)
    if campaign == "low_margin":
        _store_amounts(ctx, path, ids, MARGIN_COLS, "current_margin")
    ageing = {"status": "skipped", "detail": "client database not configured"}
    if use_client_db and available(ctx) and not reachable(ctx):
        ageing = {"status": "unavailable", "detail": "client database not reachable from here (office network / VPN needed)"}
    elif use_client_db and available(ctx):
        try:
            res = enrich_leads(ctx)
            ageing = {"status": "ok", "detail": f"{res.get('found', 0)} of {res.get('looked_up', 0)} open clients found in the client database"}
        except Exception as e:  # unreachable from the cloud, wrong password, VPN down...
            log.warning("Client DB ageing skipped: %s", e)
            ageing = {"status": "unavailable", "detail": f"client database not reachable from here ({type(e).__name__})"}
    if campaign == "reactivation":
        from .assign import release_unengaged_reactivation
        from .heat import refresh_heat

        refresh_heat(ctx)  # score first, so only the unengaged are released
        released = release_unengaged_reactivation(ctx, ids)
    # nobody new from this list goes to a GM until the SM presses Assign (clients a GM already has stay with them)
    ctx.conn.execute(f"UPDATE leads SET on_hold=1 WHERE lead_id IN ({','.join('?' * len(ids))}) AND status='open' "
                     "AND assigned_to IS NULL", tuple(ids))
    ctx.conn.commit()
    sync = jobs.sync(ctx)  # heat score, system follow-ups, overdue checks (held clients are not assigned)

    marks = ",".join("?" * len(ids))
    from .assign import engaged_enough

    leads = rows(ctx.conn, f"SELECT lead_id, status, heat_tier, assigned_to, category, campaign, on_hold FROM leads WHERE lead_id IN ({marks})", ids)
    summary = {
        "file": path.name, "kind": ingest["kind"], "list_type": LIST_TYPES[list_type][0].replace("{target}", money(target(ctx))), "clients": len(ids), "new_leads": ingest["new_leads"],
        "already_known": known_before, "ageing": ageing,
        "tiers": dict(Counter(l["heat_tier"] for l in leads if l["status"] == "open")),
        "assigned": sum(1 for l in leads if l["status"] == "open" and l["assigned_to"]),
        "on_hold": sum(1 for l in leads if l["status"] == "open" and not l["assigned_to"] and l["on_hold"]),
        "waiting": sum(1 for l in leads if l["status"] == "open" and not l["assigned_to"] and not l["on_hold"] and engaged_enough(ctx, l)),
        "watching": sum(1 for l in leads if l["status"] == "open" and not l["assigned_to"] and not l["on_hold"] and not engaged_enough(ctx, l)),
        "released_to_watching": released,
        "closed": sum(1 for l in leads if l["status"] != "open"),
        "not_in_client_db": sum(1 for l in leads if l["category"] == "Not found in client DB"),
        "assignment_run": sync.get("assignment", {}),
    }
    cur = ctx.conn.execute(
        "INSERT INTO uploads (filename, uploaded_by, uploaded_at, client_ids, summary) VALUES (?,?,?,?,?) RETURNING id",
        (path.name, uploaded_by, iso(ctx.now()), json.dumps(ids), json.dumps(summary, default=str)),
    )
    summary["upload_id"] = cur.fetchone()["id"]
    ctx.conn.commit()
    return summary


def held_count(ctx: Ctx, client_ids: list[str]) -> int:
    """Clients in this list still waiting for the SM to press Assign."""
    if not client_ids:
        return 0
    return one(ctx.conn, f"SELECT COUNT(*) AS n FROM leads WHERE on_hold=1 AND status='open' AND assigned_to IS NULL "
                         f"AND lead_id IN ({','.join('?' * len(client_ids))})", tuple(client_ids))["n"]


def assign_upload(ctx: Ctx, upload: dict, assigned_by: str | None = None) -> dict:
    """The separate Assign step: release this list's held clients and hand them to GMs (payout, 🔥, 🌤️, ❄️ first)."""
    from .assign import assign_new_leads
    from .followups import create_system_followups

    ids = upload["client_ids"]
    held = held_count(ctx, ids)
    ctx.conn.execute(f"UPDATE leads SET on_hold=0, released_at=? WHERE lead_id IN ({','.join('?' * len(ids))})",
                     (iso(ctx.now()), *ids))
    ctx.conn.commit()
    run = assign_new_leads(ctx)  # hottest first across everything waiting, so capacity goes where it matters
    create_system_followups(ctx)
    summary = {**upload["summary"], "assigned_at": iso(ctx.now()), "assigned_by": assigned_by, "assign_run": run}
    ctx.conn.execute("UPDATE uploads SET summary=? WHERE id=?", (json.dumps(summary, default=str), upload["id"]))
    ctx.conn.commit()
    return {"released": held, **run}


PAYOUT_AMOUNT_COLS = ["payout amount", "withdrawal amount", "amount", "payout", "requested amount", "payout amt"]
PAYOUT_DATE_COLS = ["payout date", "withdrawal date", "request date", "requested on", "date"]


MARGIN_COLS = ["current margin", "margin", "ledger balance", "ledger", "balance", "available margin", "funds",
               "deposit amount", "deposit", "total deposit", "first deposit"]


def _store_amounts(ctx: Ctx, path: Path, ids: list[str], columns: list[str], key: str) -> int:
    """Copy an amount column (e.g. current margin) from the uploaded file onto each client's lead."""
    from .ingest import USER_ID_COLS, _find, _read_sheets

    wanted, found = set(ids), {}
    for df in _read_sheets(path).values():
        uid, col = _find(df, "", USER_ID_COLS), _find(df, "", columns)
        if not uid or not col:
            continue
        for r in df.to_dict("records"):
            cid = norm_client_id(r[uid])
            try:
                if cid in wanted and str(r[col]).strip():
                    found[cid] = float(str(r[col]).replace(",", "").replace("₹", ""))
            except ValueError:
                pass
    for cid, amount in found.items():
        lead = one(ctx.conn, "SELECT extra FROM leads WHERE lead_id=?", (cid,))
        if lead:
            extra = json.loads(lead["extra"] or "{}")
            extra[key] = amount
            ctx.conn.execute("UPDATE leads SET extra=? WHERE lead_id=?", (json.dumps(extra), cid))
    ctx.conn.commit()
    return len(found)


def _store_payout_details(ctx: Ctx, path: Path, ids: list[str]) -> None:
    """Pick up payout amount / date columns when the file has them, for the opener ('the ₹25,000 withdrawal')."""
    from .context import parse_dt
    from .ingest import USER_ID_COLS, _find, _read_sheets

    wanted, found = set(ids), {}
    for df in _read_sheets(path).values():
        uid = _find(df, "", USER_ID_COLS)
        amt = _find(df, "", PAYOUT_AMOUNT_COLS)
        when = _find(df, "", PAYOUT_DATE_COLS)
        if not uid or not (amt or when):
            continue
        for r in df.to_dict("records"):
            cid = norm_client_id(r[uid])
            if cid not in wanted:
                continue
            d = found.setdefault(cid, {})
            try:
                if amt and str(r[amt]).strip():
                    d["payout_amount"] = float(str(r[amt]).replace(",", "").replace("₹", ""))
            except ValueError:
                pass
            dt = parse_dt(r[when], ctx.tz) if when else None
            if dt:
                d["payout_date"] = iso(dt)
    for cid, d in found.items():
        lead = one(ctx.conn, "SELECT extra FROM leads WHERE lead_id=?", (cid,))
        if lead:
            extra = json.loads(lead["extra"] or "{}")
            extra.update(d)
            ctx.conn.execute("UPDATE leads SET extra=? WHERE lead_id=?", (json.dumps(extra), cid))
    ctx.conn.commit()


def list_uploads(ctx: Ctx, limit: int = 20) -> list[dict]:
    out = []
    for u in rows(ctx.conn, "SELECT * FROM uploads ORDER BY id DESC LIMIT ?", (limit,)):
        out.append({**u, "client_ids": json.loads(u["client_ids"]), "summary": json.loads(u["summary"] or "{}")})
    return out


def batch_rows(ctx: Ctx, client_ids: list[str]) -> list[dict]:
    """Rows for just this list, in priority order, with the rank inside the list and in the whole base."""
    from .sheet_export import build_rows

    data = build_rows(ctx, lead_ids=list(client_ids))
    data.sort(key=lambda r: (r["status"] != "open", -(r["heat"] or 0)))
    for i, r in enumerate(data, 1):
        r["batch_rank"] = i if r["status"] == "open" else ""
    return data


def _sheet_view(r: dict) -> dict:
    """The merged columns of the analysed list, from one batch row."""
    assigned = bool(r["gm"]) and not r["gm"].startswith("Waiting")
    reg = f"registered {r['registered_on']} ({r['days_since_registration']} d)" if r.get("registered_on") else ""
    profile = [r.get("list"), r.get("category") if r.get("category") != r.get("list") else "", reg,
               r.get("deposit_status"), f"last trade {r['last_trade_date']}" if r.get("last_trade_date") else "",
               f"Sarthi {r['sarthi_status']}" if r.get("sarthi_status") else "",
               f"holdings ₹{r['holdings']:,}" if r.get("holdings") else ""]
    core = [p for p in (r.get("core_detail") or "").split("; ") if p][:3]
    activity = "\n".join(([f"Last active: {r['last_active_ago']}"] if r.get("last_active_ago") else []) + core + ([f"Interested in: {r['segments']}"] if r.get("segments") else [])
                         + ([f"Momentum: {r['momentum']}"] if r.get("momentum") and r["momentum"] != "Inactive" else []))
    when = " · ".join(x for x in [f"Due {r['next_due']}" if r.get("next_due") else "",
                                   f"best {r['best_time']}" if r.get("best_time") else ""] if x)
    status = r["status"] if r["status"] != "open" or assigned or not r["gm"] else f"open: {r['gm'][0].lower()}{r['gm'][1:]}"
    return {"priority_score": f"{r['priority']} · {r['heat']}", "profile": " · ".join(x for x in profile if x),
            "client_line": " · ".join(str(x) for x in [r["client"], r.get("name"), r.get("phone")] if x),
            "activity": activity or "No app activity seen", "when": when,
            "gm_assigned": r["gm"] if assigned else None, "status_line": status}


def _write_sheet(ws, data: list[dict], columns: list[tuple[str, str]], widths: dict, freeze: str) -> None:
    ws.append([h for h, _ in columns])
    for r in data:
        ws.append([r.get(k, "") for _, k in columns])
    for i, (_, k) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(k, 13)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(data) + 1, 2)}"


def _timeline_sheet(ws, data: list[dict], ctx: Ctx) -> None:
    """One row per client: recent windows, active days, peak hour, then a day-by-day and an hour-of-day heatmap."""
    from .timeline import DAYS_SHOWN, WINDOWS_H

    days = [ctx.now().date() - timedelta(days=i) for i in range(DAYS_SHOWN)]
    head = (["Rank", "Priority", "Client", "Last active"] + [f"Actions last {w} h" for w in WINDOWS_H]
            + ["Active days (7 d)", "Active days (30 d)", "Peak hour"] + [d.strftime("%a %d %b") for d in days]
            + [f"{h:02d}:00" for h in range(24)])
    ws.append(head)
    for r in data:
        t = r["timeline"]
        ws.append([r.get("batch_rank", ""), r["priority_score"], r["client"], t["last_active"]] + [t[f"last_{w}h"] for w in WINDOWS_H]
                  + [t["active_days_7"], t["active_days_30"], t["peak_hour"]] + [t["days"][d] for d in days]
                  + [t["hours"][h] for h in range(24)])
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
    for i, w in enumerate([7, 12, 16, 16] + [9] * len(WINDOWS_H) + [9, 9, 13] + [8] * DAYS_SHOWN + [6] * 24, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    n = max(len(data) + 1, 2)
    first_day = 5 + len(WINDOWS_H) + 3
    blocks = [(5, 4 + len(WINDOWS_H)), (first_day, first_day + DAYS_SHOWN - 1), (first_day + DAYS_SHOWN, first_day + DAYS_SHOWN + 23)]
    for a, b in blocks:  # white → green heatmap per block, so busy hours and days stand out
        ws.conditional_formatting.add(f"{get_column_letter(a)}2:{get_column_letter(b)}{n}",
                                      ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                     end_type="max", end_color="2E8B57"))
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{n}"


def batch_workbook(ctx: Ctx, upload: dict) -> Workbook:
    from .timeline import timeline

    data = []
    for r in batch_rows(ctx, upload["client_ids"]):
        lead = one(ctx.conn, "SELECT * FROM leads WHERE lead_id=?", (r["lead_id"],)) or {"lead_id": r["lead_id"]}
        r["timeline"] = timeline(ctx, lead)
        r["last_active_ago"] = r["timeline"]["last_active"]
        data.append({**r, **_sheet_view(r)})
    s = upload["summary"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Analysed list"
    _write_sheet(ws, data, BATCH_COLUMNS, WIDTHS, "D2")  # rank, priority and client stay in view
    keys = [k for _, k in BATCH_COLUMNS]
    for row in ws.iter_rows(min_row=2, min_col=keys.index("script_insight") + 1, max_col=keys.index("script_insight") + 1):
        row[0].font = Font(name="Arial", size=10, bold=True)  # the lines a GM reads first
    n = max(len(data) + 1, 2)
    last = get_column_letter(len(BATCH_COLUMNS))
    col = {k: get_column_letter(i) for i, k in enumerate(keys, 1)}  # formulas follow the column order
    for tier, color in TIER_FILL.items():
        ws.conditional_formatting.add(f"A2:{last}{n}", FormulaRule(formula=[f'ISNUMBER(SEARCH("{tier.split()[-1]}",${col["priority_score"]}2))'],
                                                                 fill=PatternFill("solid", fgColor=color)))
    detail_widths = {k: 24 for k in ("call_ideas_text", "objections_hi", "levers", "core_detail", "non_core_detail", "heat_reasons",
                                     "closing_text", "script_insight", "script_ask")} | {"call_ideas_text": 70, "script_text": 70}
    _timeline_sheet(wb.create_sheet("Activity timeline"), data, ctx)
    _write_sheet(wb.create_sheet("Details"), data, DETAIL_COLUMNS, detail_widths, "F2")

    sm = wb.create_sheet("Summary")
    P, G, S, C = col["priority_score"], col["gm_assigned"], col["status_line"], col["client_line"]
    sheet = "'Analysed list'!"
    count = lambda *pairs: "=COUNTIFS(" + ",".join(f'{sheet}{c}2:{c}{n},"{v}"' for c, v in pairs) + ")"
    lines = [
        ("Customer list analysis", ""),
        ("File", s.get("file", upload.get("filename"))),
        ("Uploaded", f"{upload.get('uploaded_at', '')[:16].replace('T', ' ')} by {upload.get('uploaded_by') or ''}"),
        ("", ""),
        ("Clients in the file", count((C, "<>"))),
        ("New to the system", s.get("new_leads", 0)),
        ("Already known", s.get("already_known", 0)),
        ("🔥 Hot", count((P, "🔥 Hot*"), (S, "open*"))),
        ("🌤️ Warm", count((P, "🌤️ Warm*"), (S, "open*"))),
        ("❄️ Cold", count((P, "❄️ Cold*"), (S, "open*"))),
        ("Assigned to a GM", count((G, "<>"), (S, "open"))),
        ("Not assigned yet (press Assign)", count((S, "open: waiting for assign*"))),
        ("Waiting for GM capacity", count((S, "open: waiting for capacity*"))),
        ("Re-activation: watching for engagement", count((S, "open: waiting for engagement*"))),
        ("Client database ageing", s.get("ageing", {}).get("detail", "")),
        ("", ""),
        ("How the heat score works", "core activity ×10, non-core ×3; halves every 3 days since last done; more repeats count more "
                                     "(capped); registration counts as core. 0–100: 🔥 ≥ 70, 🌤️ ≥ 40, ❄️ below."),
        ("Compliance", "Pitch ideas never recommend specific securities or promise returns."),
    ]
    for r, (a, b) in enumerate(lines, 1):
        sm.cell(row=r, column=1, value=a).font = Font(name="Arial", bold=r == 1 or not b, size=14 if r == 1 else 10)
        sm.cell(row=r, column=2, value=b).font = Font(name="Arial", size=10)
    sm.column_dimensions["A"].width = 28
    sm.column_dimensions["B"].width = 90
    wb.move_sheet("Summary", offset=-(len(wb.sheetnames) - 1))  # Summary first
    wb.active = 0
    return wb
